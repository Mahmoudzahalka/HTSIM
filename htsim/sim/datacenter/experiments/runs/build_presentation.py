#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os
try:
    from PIL import Image
    def dims(p):
        w,h=Image.open(p).size; return w/h
except Exception:
    def dims(p): return 2.6

RUNS="/home/mahmoud_murad_allaah/HTSIM/htsim/sim/datacenter/experiments/runs"
UG=os.path.join(RUNS,"uet","graphs"); IG=os.path.join(RUNS,"ib_dcqcn","graphs")
RG=os.path.join(RUNS,"rail_optimized","graphs")
OUT=os.path.join(RUNS,"IB_vs_UET_presentation.pptx")

DARK=RGBColor(0x1F,0x38,0x64); IBC=RGBColor(0x00,0x72,0xB2); UETC=RGBColor(0xD5,0x5E,0x00)
INK=RGBColor(0x22,0x22,0x22); GREY=RGBColor(0x55,0x55,0x55); WHITE=RGBColor(0xFF,0xFF,0xFF)
LIGHT=RGBColor(0xEE,0xF1,0xF7); RAILC=RGBColor(0x00,0x9E,0x73)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height
BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def rect(s,l,t,w,h,color):
    sh=s.shapes.add_shape(1,l,t,w,h); sh.fill.solid(); sh.fill.fore_color.rgb=color
    sh.line.fill.background(); return sh
def tb(s,l,t,w,h,anchor=MSO_ANCHOR.TOP):
    b=s.shapes.add_textbox(l,t,w,h); b.text_frame.word_wrap=True; b.text_frame.vertical_anchor=anchor
    return b.text_frame
def para(tf,text,size=18,color=INK,bold=False,bullet=False,level=0,space=6,align=PP_ALIGN.LEFT,italic=False):
    p=tf.paragraphs[0] if (len(tf.paragraphs)==1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.level=level; p.alignment=align; p.space_after=Pt(space)
    r=p.add_run(); r.text=("• " if bullet else "")+text
    r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.italic=italic
    r.font.name="Calibri"; return p

def header(s,title,accent=DARK):
    rect(s,0,0,SW,Inches(1.05),accent)
    tf=tb(s,Inches(0.5),Inches(0.12),Inches(12.3),Inches(0.8),MSO_ANCHOR.MIDDLE)
    para(tf,title,size=26,color=WHITE,bold=True)

def content(title,bullets,accent=DARK,sub=None):
    s=slide(); header(s,title,accent)
    top=Inches(1.35)
    if sub:
        tf=tb(s,Inches(0.6),Inches(1.15),Inches(12),Inches(0.5)); para(tf,sub,size=15,color=GREY,italic=True); top=Inches(1.7)
    tf=tb(s,Inches(0.7),top,Inches(12),Inches(5.6))
    for b in bullets:
        if isinstance(b,tuple): txt,lvl=b
        else: txt,lvl=b,0
        para(tf,txt,size=(19 if lvl==0 else 16),color=(INK if lvl==0 else GREY),
             bold=(lvl==0 and txt.endswith(":")),bullet=True,level=lvl,space=(8 if lvl==0 else 3))
    return s

def imageslide(title,img,takeaway,accent=DARK,imgw=12.0):
    s=slide(); header(s,title,accent)
    ar=dims(img); w=Inches(imgw); h=Emu(int(w/ar))
    maxh=Inches(5.05)
    if h>maxh: h=maxh; w=Emu(int(h*ar))
    left=Emu(int((SW-w)/2)); top=Inches(1.25)
    s.shapes.add_picture(img,left,top,width=w,height=h)
    ty=top+h+Inches(0.12)
    bar=rect(s,Inches(0.6),ty,Inches(12.13),Inches(0.72),LIGHT); bar.line.fill.background()
    tf=tb(s,Inches(0.8),ty+Inches(0.04),Inches(11.7),Inches(0.64),MSO_ANCHOR.MIDDLE)
    para(tf,"Takeaway: "+takeaway,size=14.5,color=DARK,bold=False)
    return s

def two_images(title,imgL,imgR,capL,capR,accent=DARK):
    s=slide(); header(s,title,accent)
    for img,cap,lx in [(imgL,capL,0.4),(imgR,capR,6.87)]:
        ar=dims(img); w=Inches(6.0); h=Emu(int(w/ar))
        if h>Inches(4.9): h=Inches(4.9); w=Emu(int(h*ar))
        s.shapes.add_picture(img,Inches(lx),Inches(1.45),width=w,height=h)
        tf=tb(s,Inches(lx),Inches(6.5),Inches(6.0),Inches(0.7)); para(tf,cap,size=13,color=GREY,align=PP_ALIGN.CENTER)
    return s

def section(title,subtitle=None,accent=DARK):
    s=slide(); rect(s,0,0,SW,SH,accent)
    tf=tb(s,Inches(1),Inches(2.9),Inches(11.3),Inches(1.4),MSO_ANCHOR.MIDDLE)
    para(tf,title,size=40,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    if subtitle:
        tf2=tb(s,Inches(1.5),Inches(4.3),Inches(10.3),Inches(1)); para(tf2,subtitle,size=18,color=RGBColor(0xCF,0xD8,0xEA),align=PP_ALIGN.CENTER)
    return s

# ---------------- 1. Title ----------------
s=slide(); rect(s,0,0,SW,SH,DARK); rect(s,0,Inches(4.55),SW,Inches(0.08),IBC)
tf=tb(s,Inches(0.9),Inches(2.0),Inches(11.5),Inches(2.2),MSO_ANCHOR.MIDDLE)
para(tf,"Lossless vs Lossy Datacenter Transports",size=40,color=WHITE,bold=True)
para(tf,"IB (RoCE + DCQCN)  vs  Ultra Ethernet (UEC)  —  flat fabrics and rail-optimized GPU clusters",size=21,color=RGBColor(0xCF,0xD8,0xEA),space=4)
tf2=tb(s,Inches(0.95),Inches(4.8),Inches(11),Inches(1.5))
para(tf2,"Two congestion philosophies, measured on a flat 200G fat-tree and on a 1024-GPU rail-optimized cluster with NVLink",size=15,color=RGBColor(0x9F,0xB2,0xD0),italic=True)

# ---------------- 2. Motivation ----------------
content("Motivation",[
 "AI / HPC datacenters are bottlenecked by the network, not compute: distributed training and storage are dominated by collective communication (incast, all-to-all, all-reduce).",
 "These patterns create extreme, bursty congestion — many senders hitting one link at once — which a naïve fabric handles badly (drops, retransmissions, collapse).",
 "Two competing industry philosophies for the fabric:",
 ("InfiniBand-style / RoCE: keep the fabric LOSSLESS (PFC back-pressure) and add DCQCN congestion control to slow senders before buffers overflow.",1),
 ("Ultra Ethernet (UEC): let the fabric DROP, but spray each packet across all paths and retransmit fast — tolerate loss instead of preventing it.",1),
 "Question: which wins, on which workloads, and what does each pay for its performance?",
 "Approach: build BOTH stacks in the htsim packet simulator and measure them under identical conditions.",
])

# ---------------- 3. Two philosophies ----------------
content("The two stacks",[
 "IB — RoCE transport + DCQCN + lossless PFC fabric:",
 ("Per-flow ECMP (one path per flow, in-order delivery).",1),
 ("PFC pauses upstream when a switch buffer fills → no packet loss.",1),
 ("DCQCN: switches mark packets (ECN) as queues build; senders cut rate on the signal.",1),
 "UET — UEC transport (NSCC congestion control) + lossy fabric:",
 ("Per-packet spray across ALL paths → uses full bisection bandwidth.",1),
 ("Packet trimming: overflowing packets are dropped (payload trimmed) and NACK'd → fast retransmit.",1),
 ("Reorder-tolerant receiver (spray causes heavy reordering).",1),
 "Both simulated in htsim on the SAME topology and workloads → a controlled, apples-to-apples comparison.",
])

# ---------------- 4. Experimental setup ----------------
content("Experimental setup — matched on both sides",[
 "Topology: 3-tier fat-tree, 200 Gbps links (a3_200g topos).",
 "Scale: 128 nodes and 1024 nodes.",
 "Oversubscription sweep: 1:1, 4:1, 8:1 (core bandwidth reduced — models cost-driven real fat-trees).",
 "Flow sizes: 16, 64, 100 MB (bulk collective-style transfers).",
 "Identical traffic matrices (.cm) fed to both stacks; identical link speed & topology.",
 "The only intentional difference besides the transport: MTU 4000 (IB) vs 4150 (UET) — each stack's native default; effect is small and proportional.",
], sub="Everything held constant except the transport/fabric — so differences are attributable to the stack.")

# ---------------- 5. Workloads part 1 ----------------
content("Workloads we ran — and what they mimic  (1/2)",[
 "Permutation — random & full-bisection:  each node sends to one other node.",
 ("Mimics uniform shuffle / general point-to-point traffic; full-bisection maximally stresses the network core.",1),
 "Incast (random) — many senders → one receiver (all-to-one):",
 ("The classic datacenter pain point: distributed-storage reads, MapReduce shuffle-to-reducer, parameter-server gather, partition/aggregate query fan-in.",1),
 "Incast (remote) — incast where senders are the far half of the fabric:",
 ("Same fan-in but forced across the core → stresses incast AND core bandwidth together.",1),
 "Outcast-incast — combined fan-out then fan-in:",
 ("Scatter-then-gather / tree-reduction communication.",1),
], accent=UETC)

# ---------------- 6. Workloads part 2 (collectives) ----------------
content("Workloads we ran — and what they mimic  (2/2)",[
 "All-to-all (a2a) — every node sends to every other node:",
 ("The most demanding collective. Mimics MPI_Alltoall: Mixture-of-Experts routing, embedding/table shuffles in recommendation models, FFT transposes, distributed sort.",1),
 ("Serial vs concurrent variants: staged rounds vs all pairs firing at once (peak stress).",1),
 "Ring all-reduce — ring-structured gradient aggregation:",
 ("The dominant collective in data-parallel deep-learning training (gradient sync every step).",1),
 "Oversubscription (1:1 → 8:1) mimics real fat-trees thinned at the core for cost — testing each stack as the network gets more contended.",
 "Together these span the traffic that actually limits AI/HPC clusters — from benign permutation to pathological all-to-one incast.",
], accent=UETC)

# ---------------- 7. What we measure ----------------
content("What we measure",[
 "Common (directly comparable between stacks):",
 ("Makespan — time for the whole workload to complete (the bottom line).",1),
 ("Flow-completion time (FCT) & slowdown = FCT ÷ ideal (tail latency).",1),
 ("Per-tier link utilization (how well each fabric uses the network).",1),
 ("Completion / robustness — did every flow actually finish?",1),
 "Each stack's native congestion signal (different mechanisms):",
 ("UET: packet-trim / NACK rate and retransmission count (its cost of dropping).",1),
 ("IB: PFC pause count & pause-time, and peak buffer occupancy (its cost of back-pressure).",1),
 ("IB also: per-flow Jain fairness (how evenly bandwidth is shared).",1),
])

# ================= UET RESULTS =================
section("Results — UET (Ultra Ethernet)","Lossy fabric, per-packet spray + trimming, NSCC",UETC)
imageslide("UET — makespan vs oversubscription (128 nodes)",os.path.join(UG,"fig1_makespan_vs_os_128.png"),
 "Spray uses all paths, so permutation finishes fast; incast/all-to-all dominate and grow steeply with oversubscription.",UETC)
imageslide("UET — congestion (packet-trim / NACK rate)",os.path.join(UG,"fig2_nack_vs_os_128.png"),
 "Trim rate climbs sharply with oversubscription — under 8:1 a large fraction of packets are dropped and retransmitted (up to billions of retransmits for all-to-all).",UETC)
imageslide("UET — per-tier link utilization (128 nodes, 100 MB)",os.path.join(UG,"fig3_utilization_heatmaps_128.png"),
 "Per-packet spray saturates links to ~100% across tiers — UET extracts maximum bandwidth from the fabric.",UETC)
imageslide("UET — serial vs concurrent all-to-all (128 nodes)",os.path.join(UG,"fig4_a2a_serial_vs_concurrent_128.png"),
 "Concurrent all-to-all is far harder than staged/serial — peak contention drives both makespan and trim rate up.",UETC)

# ================= IB RESULTS =================
section("Results — IB (RoCE + DCQCN)","Lossless PFC fabric, per-flow ECMP, DCQCN",IBC)
imageslide("IB — makespan vs oversubscription (128 nodes)",os.path.join(IG,"fig1_makespan_vs_os_128.png"),
 "Incast is flat in OS (receiver is the bottleneck); permutation grows with OS because IB uses one path per flow (no spray).",IBC)
imageslide("IB — makespan vs oversubscription (1024 nodes)",os.path.join(IG,"fig2_makespan_vs_os_1024.png"),
 "Same structure holds at scale; IB completes every 1024-node workload, including the extreme incasts.",IBC)
imageslide("IB — PFC back-pressure (its congestion cost)",os.path.join(IG,"fig5_pfc_pause_vs_os.png"),
 "IB is lossless (zero retransmits) — instead it pays in PFC pause-time, which rises steeply with oversubscription, most for incast.",IBC)
imageslide("IB — per-tier link utilization (128 nodes, 100 MB)",os.path.join(IG,"fig4_utilization_heatmaps_128.png"),
 "DCQCN paces senders → link utilization is lower and smoother than UET's saturation (bandwidth traded for gentleness).",IBC)
two_images("IB — fairness & tail slowdown  (unique to IB instrumentation)",
 os.path.join(IG,"fig6_fairness_vs_os.png"),os.path.join(IG,"fig7_slowdown_p99_vs_os.png"),
 "Jain fairness: incast is near-perfectly fair; permutation & outcast grow UNFAIR under oversubscription (down to ~0.4).",
 "Tail slowdown (p99): incast pays a huge tail (100×+), permutation stays low.",IBC)
imageslide("IB — peak PFC buffer occupancy (lossless buffer cost)",os.path.join(IG,"fig8_peak_buffer_vs_os.png"),
 "How much buffer the lossless fabric actually needs — the memory price of never dropping a packet.",IBC)

# ================= ANALYSIS =================
section("Analysis — putting them side by side")
content("Analysis (1): speed",[
 "UET is generally faster — makespan ratio IB/UET ≈ 2.2× on average; IB is faster on only ~18 of 96 matched runs.",
 "Biggest gap is PERMUTATION: IB is 3–4× slower.",
 ("Why: UET sprays every packet across all 128 paths and saturates the fabric; IB pins each flow to one ECMP path (RoCE needs near-in-order delivery), so flows collide.",1),
 "IB wins the HARD incasts: on 1024-node random incast IB is ~12% faster with a better tail.",
 ("Why: the receiver port is the physical bottleneck — spray can't help — and IB avoids UET's retransmission overhead there.",1),
], accent=DARK)
content("Analysis (2): robustness & the cost of speed",[
 "Robustness — IB completed ALL 99 runs. UET left the most extreme cases unfinished:",
 ("1024-node random incast at 8:1 oversubscription — incomplete (all sizes).",1),
 ("1024-node all-to-all — timed out.",1),
 "The cost each stack pays under heavy congestion is opposite:",
 ("UET burns BANDWIDTH on retransmissions — up to ~3 billion retransmits in one all-to-all run (each packet sent ~8× at 8:1 oversubscription).",1),
 ("IB burns TIME on back-pressure — zero retransmits, but PFC pauses stall senders (hundreds of port-seconds), and it needs large lossless buffers.",1),
 "Utilization mirrors this: UET saturates links (~100%); IB deliberately runs them cooler (DCQCN pacing).",
], accent=DARK)


# ================= PART 2: RAIL-OPTIMIZED =================
section("Part 2 — Rail-Optimized Clusters + NVLink","Modelling how AI clusters are actually built",RAILC)

content("Why rail-optimized? (the real hardware)",[
 "Everything so far assumed a FLAT fabric: every GPU is just an endpoint on a fat-tree. Real AI clusters are not built that way.",
 "A DGX-class server holds 8 GPUs joined internally by NVLink/NVSwitch — roughly 450 GB/s per GPU per direction, ~18x a 200G NIC.",
 "Each GPU also has its OWN 200-400G NIC, and each of those NICs is wired to a DIFFERENT switch — its 'rail'. 8 GPUs per server => 8 rails.",
 "So a leaf switch does NOT hold whole servers: it holds ONE GPU from each of many servers — that is what makes it a rail.",
 "Consequence: two GPUs on the same rail are one hop apart; two GPUs in the same server are NOT (they are on different rails) — but they have NVLink between them.",
 "Purpose: make the traffic that collectives generate either stay inside the server (NVLink, very fast) or stay inside a rail (one leaf, no spine).",
], accent=RAILC)

content("The topology we built (1024 GPUs, DGX SuperPOD-class)",[
 "128 servers x 8 GPUs = 1024 GPUs; each GPU has its own 200G NIC on its own rail.",
 "8 rails; each rail is served by 4 leaf switches (32 servers each) => 32 leaves total, each 32 down + 32 up = 64 ports (real Quantum-2-class radix).",
 "32 spines, 2-tier leaf-spine, fully NON-BLOCKING, diameter 4.",
 "Host->leaf mapping (this is the whole difference):",
 ("GPU g  ->  rail r = g mod 8,  server s = g / 8,  leaf = r*4 + s/32",1),
 ("i.e. hosts are attached STRIDED, not consecutively. A leaf holds GPU r of 32 different servers.",1),
 "Control topology: identical shape (32 leaves, 32 spines, same speeds) but CONSECUTIVE attachment — so any difference is attributable to the rail assignment alone, not the switch layout.",
 "A second topology models the NVLink domain: 128 leaves = 128 servers, 8 GPUs each at 3600 Gbps, 100 ns links.",
], accent=RAILC)

content("How NVLink is modelled — and why it was essential",[
 "A GPU has TWO attachments: NVLink to its 7 server-mates, and a rail NIC to everything else. A fat-tree gives each host only one, so we run two topologies.",
 "The choice is made per FLOW at connect time (a flow's destination is fixed):",
 ("same server -> NVLink topology at NVLink rate;  otherwise -> rail topology at 200G.",1),
 "The fabric alone was not enough: the HOST send-rate is the binding constraint. Same NVLink path, 64 MB flow:",
 ("200G host rate -> 2601 us    vs    3600G host rate -> 145 us   (18x)",1),
 "Validated on both stacks: intra-server flows hit 145 us; cross-server flows are unchanged (never mis-routed onto NVLink); and a GPU can drive NVLink and its rail NIC simultaneously.",
 "Ablation finding: rails WITHOUT NVLink are actively harmful (IB 2.2x worse than flat) — because rails scatter a server's GPUs across 8 leaves. Rails and NVLink only make sense together.",
], accent=RAILC)

content("The workloads: rail-AWARE collectives",[
 "A rail fabric only pays off if the collective is written for it. Standard collectives are rail-agnostic and would just pay the cost (7/8 of destinations are cross-rail).",
 "Rail-aware all-reduce (hierarchical, as used on DGX):",
 ("1. intra-server reduce-scatter over NVLink  2. rail-local all-reduce (recursive halving/doubling)  3. intra-server all-gather over NVLink",1),
 "Rail-aware all-to-all = the MoE dispatch pattern (Mixture-of-Experts routing):",
 ("1. NVLink shuffle so data for rail r sits on the GPU at rail r   2. rail-local all-to-all",1),
 ("trades ~1.87x more bytes for perfect locality — the bytes are cheap because they ride NVLink",1),
 "Both generators assert the key property: every flow is either intra-server or same-rail — 0 flows cross a spine.",
], accent=RAILC)

imageslide("Rail results — makespan vs the shape-matched control",os.path.join(RG,"fig1_rail_vs_flat.png"),
 "Same switch layout in both; only the host->leaf assignment differs. IB (blue) gains on every configuration; UET (orange) barely moves.",RAILC)
imageslide("Rail benefit, and the IB/UET ratio",os.path.join(RG,"fig2_rail_benefit_and_ratio.png"),
 "Left: rails help IB by 7-33% on all six runs, UET within noise. Right: on MoE all-to-all the stacks converge — and at 100 MB IB is 11% FASTER.",RAILC)
imageslide("What each stack pays on the rail fabric",os.path.join(RG,"fig3_rail_congestion_cost.png"),
 "Two different currencies: IB burns time in PFC back-pressure (up to ~9.7 s of pause), UET burns bandwidth in retransmissions (up to 3.8 M packets).",RAILC)

content("Analysis — three findings",[
 "1. Rail-optimization is an IB optimization.",
 ("Rails helped IB on all six runs (-7% to -33%); UET stayed within noise (+9% to -8%).",1),
 ("Why: rails are a LOCALITY optimization. IB pins each flow to one path, so keeping traffic on one leaf matters. UET sprays across all paths and has spare bandwidth on a non-blocking fabric, so it is largely indifferent to where traffic goes.",1),
 "2. On MoE all-to-all the two stacks converge — IB even wins at 100 MB (0.89x).",
 ("Versus UET's usual 3-4x lead on all-reduce. This is the sharpest reversal in the study.",1),
 "3. The mechanism: rail locality does NOT reduce UET's retransmits (~0% change vs flat).",
 ("MoE all-to-all makes every GPU receive from its 127 rail peers, so the bottleneck is receiver-side INCAST at the endpoint — not congestion in the fabric.",1),
 ("An endpoint bottleneck cannot be fixed by locality, and cannot be fixed by spray either. That single fact explains all three findings — and why IB's losslessness edges ahead where UET's retransmissions become pure overhead.",1),
], accent=RAILC)

# ================= CONCLUSIONS =================
content("Conclusions",[
 "Part 1 — flat fabric: a clean speed-vs-robustness trade-off.",
 ("UET (lossy spray) is 2-4x faster on well-behaved/permutation traffic and saturates links, but retransmits enormously under congestion and fails to complete the most extreme incasts.",1),
 ("IB (lossless PFC + DCQCN) is slower but never drops a packet, completes every workload UET cannot, and has a better tail on the heavy 1024-node incasts — paying in back-pressure and buffering.",1),
 "Part 2 — rail-optimized + NVLink: the picture changes.",
 ("Rails+NVLink must be adopted together: rails alone are 2.2x WORSE than a flat fabric, because they scatter each server's GPUs across 8 leaves.",1),
 ("Rail-optimization is an IB optimization (-7% to -33% on every run); UET is indifferent, because spray already makes it insensitive to locality.",1),
 ("On MoE all-to-all the two stacks converge and IB overtakes UET at 100 MB — the workload is bound by receiver-side incast at the endpoint, which neither locality nor spray can relieve.",1),
 "Overall: which transport wins is workload- and topology-dependent, not absolute. UET's advantage comes from path diversity, so it shrinks exactly where the bottleneck stops being the fabric.",
], accent=DARK)

content("Future work",[
 "Confirm the endpoint-incast explanation directly by instrumenting per-receiver backlog (currently inferred from retransmit behaviour, not measured).",
 "Tune DCQCN (ECN threshold K, min rate) — K=8 was one operating point, never swept.",
 "Adaptive routing for IB (ecmp_ar): would give IB some path diversity and likely close part of the permutation gap.",
 "Rail-aware workloads beyond these two: all-gather / reduce-scatter (FSDP, ZeRO-3) and pipeline-parallel point-to-point.",
 "Sensitivity to the NVLink rate, and to rail count / servers-per-leaf (the sweep fixed 8 rails x 4 leaves).",
 "Larger scale (8192 GPUs) and real AI communication traces (GOAL / ATLAHS).",
 "Symmetric fairness metrics: per-flow fairness for trigger-staged workloads needs per-flow start times (currently an artifact for staged collectives).",
], accent=DARK)

# closing
s=slide(); rect(s,0,0,SW,SH,DARK)
tf=tb(s,Inches(1),Inches(3.1),Inches(11.3),Inches(1.3),MSO_ANCHOR.MIDDLE)
para(tf,"Thank you",size=40,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
para(tf,"IB (RoCE/DCQCN) vs Ultra Ethernet — controlled comparison in htsim",size=18,color=RGBColor(0xCF,0xD8,0xEA),align=PP_ALIGN.CENTER)

prs.save(OUT)
print("wrote",OUT,"—",len(prs.slides._sldIdLst),"slides")
